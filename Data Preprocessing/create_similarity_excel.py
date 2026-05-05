import os
import random
import torch
import torch.nn.functional as F
import torchaudio
import soundfile as sf  # Thư viện thay thế torchaudio.load để tránh lỗi backend
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from model_arch import load_model, compute_fbank

# --- CẤU HÌNH ---
DATASET_ROOT = r"D:\Viettel\dataset\ViVoice34\VoxVietnamese_dataset\ViVoice34" 
OUTPUT_FILE = "speaker_duplicates_check.xlsx"
CONFIG_PATH = os.path.join("models", "config_34.yaml")
MODEL_PATH  = os.path.join("models", "best_model_34_new_dataset.ckpt")

SIMILARITY_THRESHOLD = 0.7
SAMPLES_PER_SPEAKER = 1 

def get_embedding_tensor(model, wav_path, device):
    """Tính embedding cho 1 file audio dùng soundfile để tránh lỗi backend"""
    try:
        # 1. Đọc file bằng soundfile (Đọc thô thành numpy array)
        data, sr = sf.read(wav_path)
        
        # 2. Chuyển sang torch.Tensor [Channel, Time]
        wav = torch.FloatTensor(data)
        if wav.ndim == 1:
            wav = wav.unsqueeze(0) # [1, Time]
        else:
            wav = wav.T # Nếu là stereo [Time, Channel] -> [Channel, Time]
        
        # 3. Resample nếu không phải 16kHz
        if sr != 16000:
            resampler = torchaudio.transforms.Resample(sr, 16000)
            wav = resampler(wav)
        
        # 4. Chuyển sang mono nếu nhiều channel
        if wav.shape[0] > 1:
            wav = torch.mean(wav, dim=0, keepdim=True)
            
        # 5. Trích xuất đặc trưng và embedding
        feats = compute_fbank(wav).to(device)
        with torch.no_grad():
            emb = model(feats)
            if isinstance(emb, tuple):
                emb = emb[-1]
        
        emb = F.normalize(emb, p=2, dim=1)
        return emb.squeeze()
    except Exception as e:
        print(f"  ⚠️ Lỗi file {os.path.basename(wav_path)}: {e}")
        return None

def get_single_speaker_centroid(speaker_path, model, device):
    """Lấy ngẫu nhiên 1 file .wav trong folder để làm đại diện"""
    wav_files = []
    for root, _, files in os.walk(speaker_path):
        for f in files:
            if f.endswith('.wav'):
                wav_files.append(os.path.join(root, f))
    
    if not wav_files:
        return None

    selected_file = random.choice(wav_files)
    centroid = get_embedding_tensor(model, selected_file, device)
    return centroid

def save_duplicates_to_excel(duplicate_pairs, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Speaker Duplicates"
    headers = ["Speaker A", "Speaker B", "Similarity Score", "Note"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for row_idx, (spk_a, spk_b, score) in enumerate(duplicate_pairs, 2):
        ws.cell(row_idx, 1, spk_a)
        ws.cell(row_idx, 2, spk_b)
        score_cell = ws.cell(row_idx, 3, score)
        score_cell.number_format = '0.0000'
        if score > 0.8:
            score_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            ws.cell(row_idx, 4, "Trùng lặp rất cao (> 80%)")
        else:
            ws.cell(row_idx, 4, "Trùng lặp cao (> 70%)")
    
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 30
    wb.save(output_path)

def main():
    device = "cuda" if torch.cuda.is_available() else "cpu"
    print(f"🔧 Thiết bị: {device}")
    
    if not os.path.exists(CONFIG_PATH) or not os.path.exists(MODEL_PATH):
        print(f"❌ Không tìm thấy model tại {MODEL_PATH}")
        return

    model = load_model(CONFIG_PATH, MODEL_PATH, device)
    model.eval()

    # 1. Thu thập danh sách speaker CHỈ TỪ THƯ MỤC TRAIN
    speaker_paths_dict = {}
    for split in ["train"]: # <-- Đã sửa ở đây: Xóa "val" khỏi list
        split_dir = os.path.join(DATASET_ROOT, split)
        if os.path.exists(split_dir):
            for spk_name in os.listdir(split_dir):
                path = os.path.join(split_dir, spk_name)
                if os.path.isdir(path):
                    # Bỏ tiền tố split_ vì giờ chỉ có mỗi "train"
                    speaker_paths_dict[spk_name] = path

    print(f"📂 Tìm thấy {len(speaker_paths_dict)} speaker trong thư mục train.")

    # 2. Trích xuất embedding
    speaker_centroids = {}
    for idx, (name, path) in enumerate(speaker_paths_dict.items()):
        emb = get_single_speaker_centroid(path, model, device)
        if emb is not None:
            speaker_centroids[name] = emb
        if idx % 100 == 0 and idx > 0:
            print(f"   Đã xử lý {idx}/{len(speaker_paths_dict)}...")

    if len(speaker_centroids) < 2:
        print("⚠️ Không đủ dữ liệu so sánh.")
        return

    # 3. Tính toán độ tương đồng chéo
    print("🔄 Đang so khớp chéo toàn bộ hệ thống...")
    names = list(speaker_centroids.keys())
    embeddings = torch.stack([speaker_centroids[n] for n in names]).to(device)
    sim_matrix = torch.mm(embeddings, embeddings.t())

    duplicate_pairs = []
    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            score = sim_matrix[i, j].item()
            if score > SIMILARITY_THRESHOLD:
                duplicate_pairs.append((names[i], names[j], score))

    # 4. Xuất kết quả
    if duplicate_pairs:
        duplicate_pairs.sort(key=lambda x: x[2], reverse=True)
        save_duplicates_to_excel(duplicate_pairs, OUTPUT_FILE)
        print(f"✅ Xong! Tìm thấy {len(duplicate_pairs)} cặp trùng lặp. Kết quả lưu tại: {OUTPUT_FILE}")
    else:
        print("✅ Tuyệt vời! Không có speaker nào trùng lặp.")

if __name__ == "__main__":
    main()